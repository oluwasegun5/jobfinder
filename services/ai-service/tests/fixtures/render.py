"""Renders synthetic CVs to real PDF / DOCX bytes so tests exercise the actual extractors."""

import io
import textwrap
from dataclasses import dataclass

from docx import Document


@dataclass(frozen=True, slots=True)
class Table:
    """A one-row table whose cells hold paragraphs (the two-column CV layout)."""

    cells: tuple[tuple[str, ...], ...]


Block = str | Table


def _flatten(blocks: list[Block]) -> list[str]:
    lines: list[str] = []
    for block in blocks:
        if isinstance(block, Table):
            for cell in block.cells:
                lines.extend(cell)
        else:
            lines.append(block)
    return lines


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def render_pdf(blocks: list[Block], *, lines_per_page: int = 55) -> bytes:
    """A minimal text-only PDF (Helvetica, ASCII). Long lines are wrapped to fit the page."""
    lines: list[str] = []
    for line in _flatten(blocks):
        lines.extend(textwrap.wrap(line, 88, subsequent_indent="  ") or [""])
    pages = [lines[i : i + lines_per_page] for i in range(0, len(lines), lines_per_page)] or [[]]

    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{p} 0 R' for p in page_ids)}] "
        f"/Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for page_id, page_lines in zip(page_ids, pages, strict=True):
        body = "BT /F1 10 Tf 50 750 Td 13 TL\n" + "".join(
            f"({_escape(line)}) '\n" for line in page_lines
        )
        stream = (body + "ET").encode("ascii")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {page_id + 1} 0 R "
            f"/Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return bytes(out)


def render_docx(blocks: list[Block]) -> bytes:
    document = Document()
    for block in blocks:
        if isinstance(block, Table):
            table = document.add_table(rows=1, cols=len(block.cells))
            for cell, paragraphs in zip(table.rows[0].cells, block.cells, strict=True):
                cell.text = paragraphs[0]
                for paragraph in paragraphs[1:]:
                    cell.add_paragraph(paragraph)
        else:
            document.add_paragraph(block)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
